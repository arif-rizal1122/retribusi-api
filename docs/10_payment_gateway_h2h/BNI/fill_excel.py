import openpyxl

file_path = "List_Assessment_Perusahaan_pihak_ketiga_Template_Diisi_Oleh_Pihak.xlsx"
wb = openpyxl.load_workbook(file_path)
ws = wb.worksheets[0]

def set_val(r, c, val):
    cell = ws.cell(row=r, column=c)
    for merged_range in ws.merged_cells.ranges:
        if cell.coordinate in merged_range:
            min_col, min_row, max_col, max_row = merged_range.bounds
            ws.cell(row=min_row, column=min_col).value = val
            return
    cell.value = val

for row in range(4, 172):
    deskripsi = str(ws.cell(row=row, column=2).value or "")
    
    if "Nama perusahaan" in deskripsi:
        set_val(row, 5, "Bapenda Kota Baubau (M-PAD)")
    elif "Website perusahaan" in deskripsi:
        set_val(row, 5, "https://bapenda.baubaukota.go.id")
    elif "Website/ layanan yang digunakan" in deskripsi:
        set_val(row, 5, "https://api.mpad.online")
    elif "Tanggal berdiri" in deskripsi:
        set_val(row, 5, "Pemerintah Daerah")
    elif "Jenis Bisnis" in deskripsi:
        set_val(row+5, 5, "V")
        set_val(row+5, 6, "Pemerintahan Daerah (Penerimaan Pajak Daerah)")
    elif "Struktur organisasi" in deskripsi:
        set_val(row, 5, "Terlampir")
    elif "Penanggung jawab perusahaan" in deskripsi:
        set_val(row, 5, "Kepala Bapenda Kota Baubau")
    elif "terdaftar di OJK / BI" in deskripsi:
        set_val(row+1, 5, "V")
        set_val(row+1, 6, "Pemerintahan Daerah tidak terdaftar di OJK/BI")
    elif "sertifikasi terkait IT" in deskripsi:
        set_val(row+3, 5, "V")
        set_val(row+3, 6, "Mengikuti Standar Keamanan SPBE Pemerintah Daerah")
    elif "Lokasi server dan/atau data" in deskripsi:
        set_val(row+3, 5, "V")
    elif "credential data pada perusahaan" in deskripsi:
        set_val(row, 5, "V")
    elif "penanggung jawab keamanan data" in deskripsi:
        set_val(row, 5, "Kepala Bidang TIK Diskominfo / Tim IT Bapenda")
    elif "jenis data credential" in deskripsi:
        set_val(row, 5, "Data Wajib Pajak (NIK, Nama, Nomor HP), Data Tagihan")
    elif "data BNI yang berada pada perusahaan" in deskripsi:
        set_val(row, 5, "Virtual Account Number, Client ID, Client Secret, Callback Signature")
    elif "klasifikasi data" in deskripsi:
        set_val(row, 5, "V")
    elif "menjaga kerahasiaan data credential" in deskripsi:
        set_val(row, 5, "V")
        set_val(row+2, 5, "V")
    elif "disposal data bank" in deskripsi:
        set_val(row+1, 5, "V")
        set_val(row+1, 6, "Sesuai regulasi, data riwayat pajak disimpan negara")
    elif "pengamanan data perusahaan anda dengan vendor 3rd party" in deskripsi:
        set_val(row, 5, "V")
    elif "lokasi data center" in deskripsi:
        set_val(row, 5, "Jakarta (Biznet Neo Cloud VPS)")
    elif "Pengelolaan data center" in deskripsi:
        set_val(row+4, 5, "V")
    elif "maka sebutkan nama vendornya" in deskripsi:
        set_val(row, 5, "PT Supra Primatama Nusantara (Biznet Neo)")
    elif "mekanisme backup data" in deskripsi:
        set_val(row+1, 5, "V")
    elif "pengujian DRC" in deskripsi:
        set_val(row+1, 5, "V")
    elif "Security Operation Center (SOC)" in deskripsi:
        set_val(row+2, 5, "V")
    elif "pengembangan aplikasi" in deskripsi:
        set_val(row+3, 5, "V")
    elif "secure coding" in deskripsi:
        set_val(row, 5, "V")
    elif "penetration testing (pentest)" in deskripsi:
        set_val(row+1, 5, "V")
    elif "vulnerability assessment" in deskripsi:
        set_val(row+1, 5, "V")
    elif "platform aplikasi" in deskripsi:
        set_val(row+5, 5, "V")
        set_val(row+5, 6, "PHP (Laravel 11), MySQL/PostgreSQL")
    elif "arsitektur aplikasi" in deskripsi:
        set_val(row, 5, "Monolithic / MVC. (Topologi Terlampir)")
    elif "proses rekonsiliasi" in deskripsi:
        set_val(row, 5, "Otomatis melalui H2H Callback SNAP BI.")
    elif "topologi jaringan" in deskripsi:
        set_val(row, 5, "Internet Publik HTTPS/TLS -> WAF -> VPS Server. (Terlampir)")
    elif "fraud detection system" in deskripsi:
        set_val(row+1, 5, "V")
    elif "manajemen risiko" in deskripsi:
        set_val(row+1, 5, "V")
    elif "flow proses bisnis" in deskripsi:
        set_val(row, 5, "Terlampir di Zip")
    elif "Jumlah endpoint yang digunakan" in deskripsi:
        set_val(row+2, 5, "1")
    elif "PC ditempatkan di lingkungan yang aman" in deskripsi:
        set_val(row, 5, "V")
    elif "login ke operating system menggunakan non-privileged user" in deskripsi:
        set_val(row, 5, "V")
    elif "Aplikasi yang di install di end point merupakan aplikasi yang legal" in deskripsi:
        set_val(row, 5, "V")
    elif "Tersedia antivirus dan active scanning" in deskripsi:
        set_val(row, 5, "V")
    elif "authentication, authorization, and accounting system yang terpusat" in deskripsi:
        set_val(row, 5, "V")
    elif "Penutupan port USB" in deskripsi:
        set_val(row, 5, "V")
    elif "Pembatasan akses internet" in deskripsi:
        set_val(row, 5, "V")
    elif "prosedur insiden response" in deskripsi:
        set_val(row+1, 5, "V")
    elif "Hotline helpdesk" in deskripsi:
        set_val(row, 5, "+62")
    elif "Lokasi helpdesk" in deskripsi:
        set_val(row, 5, "Kantor Bapenda Kota Baubau")
    elif "layanan heldpesk perusahaan" in deskripsi:
        set_val(row+1, 5, "V")
    elif "pengalaman kerja sama dengan bank" in deskripsi:
        set_val(row, 5, "V")
    elif "Sebutkan bank / lembaga keuangan apa yang telah melakukan bekerja sama" in deskripsi:
        set_val(row, 5, "Bank Sultra (BPD), Bank Rakyat Indonesia (BRI)")
    elif "pengalaman perusahaan dalam kerjasama" in deskripsi:
        set_val(row, 5, "Integrasi Host-to-Host (H2H) PBB-P2, BPHTB, 9 Pajak (Standar SNAP BI)")
    elif "security awareness" in deskripsi:
        set_val(row+1, 5, "V")

wb.save("Assessment_BNI_Terisi.xlsx")
print("Saved filled excel.")
